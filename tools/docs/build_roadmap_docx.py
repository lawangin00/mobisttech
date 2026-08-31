"""Render the canonical structural roadmap to Word; verify all visible text stays identical."""
from pathlib import Path
import hashlib
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/PROJECT_IMPLEMENTATION_ROADMAP.md'
OUTPUT = SOURCE.with_suffix('.docx')


def content_blocks(text):
    for block in re.split(r'\n\s*\n', text.strip()):
        block = block.strip()
        if block.startswith('#'):
            match = re.match(r'^(#{1,3}) (.*)$', block)
            if not match:
                raise ValueError('Unsupported heading')
            yield len(match[1]), match[2]
        else:
            if block.startswith(('- ', '| ', '```')):
                raise ValueError('Add explicit list/table/code rendering before using this form')
            yield 0, ' '.join(block.splitlines())


def set_style(doc, name, size, color, before, after, bold=False, line=1.25):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True
    fmt.keep_with_next = name.startswith('Heading') or name in ('Title', 'Subtitle')
    return style


def main():
    text = SOURCE.read_text(encoding='utf-8')
    blocks = list(content_blocks(text))
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)

    # compact_reference_guide, memo_masthead without a decorative rule.
    set_style(doc, 'Normal', 11, '000000', 0, 6)
    set_style(doc, 'Heading 1', 16, '2E74B5', 18, 10, True)
    set_style(doc, 'Heading 2', 13, '2E74B5', 14, 7, True)
    set_style(doc, 'Heading 3', 12, '1F4D78', 10, 5, True)
    # Named masthead/page-furniture overrides.
    set_style(doc, 'Title', 23, '0B2545', 0, 8, True)
    set_style(doc, 'Subtitle', 11, '555555', 0, 6)
    set_style(doc, 'Header', 9, '555555', 0, 0, line=1)
    set_style(doc, 'Footer', 9, '555555', 0, 0, line=1)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)

    section.header.paragraphs[0].text = 'mobiST Tech | Implementation roadmap'
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    for level, value in blocks:
        style = {0: 'Normal', 1: 'Title', 2: 'Heading 1', 3: 'Heading 2'}[level]
        paragraph = doc.add_paragraph(value, style)
        if level == 2 and ((re.match(r'MT-\d+ - ', value) and value != 'MT-0 - Project initialization') or value == 'HOLD / Deferred register'):
            paragraph.paragraph_format.page_break_before = True
        if value.startswith('Dependencies:'):
            paragraph.paragraph_format.keep_with_next = True
        if value.startswith(('Scope:', 'Acceptance:', 'Stage exit:')):
            label, rest = value.split(':', 1)
            paragraph.clear()
            paragraph.add_run(label + ':').bold = True
            paragraph.add_run(rest)

    doc.core_properties.title = 'mobiST Tech - Project Implementation Roadmap'
    doc.core_properties.author = 'mobiST Technologies'
    doc.core_properties.subject = 'Canonical Markdown roadmap mirror'
    doc.core_properties.comments = 'Markdown SHA-256: ' + hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    doc.save(OUTPUT)
    actual = [p.text for p in Document(OUTPUT).paragraphs if p.text]
    expected = [value for _, value in blocks]
    if actual != expected:
        raise RuntimeError('Markdown and Word content differ')
    ids = re.findall(r'^### ((?:MT-\d+\.\d+|FINAL-AUDIT)) - ', text, re.M)
    assert len(ids) == len(set(ids)), 'Duplicate roadmap IDs'
    assert ids[-1] == 'FINAL-AUDIT'
    assert sum(value.startswith('Dependencies:') for _, value in blocks) == len(ids), 'Each roadmap point must retain one Dependencies line'
    assert not any(value.startswith(('Status:', 'Stage status:')) for _, value in blocks), 'Live status belongs in PROJECT_IMPLEMENTATION_STATUS.md, not the roadmap'
    print(f'Word parity verified: {len(ids)} points, {len(blocks)} content blocks')
    print('Roadmap SHA-256: ' + hashlib.sha256(SOURCE.read_bytes()).hexdigest())


if __name__ == '__main__':
    main()
